"""Escritura del reporte en Excel (y CSV de respaldo)."""
import csv
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .util import aviso

COLUMNAS_FUNCIONES = [
    ("evento", "Evento", 38), ("fecha", "Fecha", 11), ("dia", "Día", 10), ("hora", "Hora", 9),
    ("lugar", "Lugar", 32), ("estado", "Estado", 13), ("agotado", "Agotado", 9),
    ("capacidad", "Capacidad numerada (entradas)", 12), ("vendidas", "Vendidas", 10),
    ("no_vendidas", "No vendidas", 10), ("pct_vendido", "% vendido (por entradas)", 11),
    ("recaudacion_vendida", "Recaudación vendida ($, precio lista)", 16),
    ("recaudacion_no_vendida", "Recaudación no vendida ($, precio lista)", 16),
    ("recaudacion_potencial", "Recaudación potencial ($)", 16),
    ("sectores_numerados", "Sectores numerados", 10), ("sectores_numerados_agotados", "Sectores agotados", 10),
    ("sectores_generales", "Sectores sin asiento (cancha/general)", 40),
    ("aforo_recinto", "Aforo recinto", 10), ("cancha_estimada", "Cancha estimada", 10),
    ("ocupacion_total_estimada", "% ocupación total (estim.)", 12),
    ("calidad", "Calidad del dato", 30), ("sectores_sin_dato", "Sectores numerados sin dato", 30),
    ("sectores_sin_precio", "Sectores sin precio", 20), ("estado_landing", "Estado en landing", 12),
    ("fecha_scraping", "Fecha scraping", 16), ("url_evento", "URL evento", 40), ("url_compra", "URL compra", 40),
]

COLUMNAS_SECTORES = [
    ("evento", "Evento", 34), ("fecha", "Fecha", 11), ("hora", "Hora", 8), ("lugar", "Lugar", 30),
    ("sector", "Sector", 26), ("sector_id", "Sector ID", 14), ("tipo", "Tipo", 10),
    ("precio_num", "Precio ($)", 11), ("agotado", "Agotado", 9),
    ("capacidad", "Capacidad (entradas)", 11), ("vendidas", "Vendidas", 10), ("no_vendidas", "No vendidas", 10),
    ("pct_vendido", "% vendido", 10), ("recaudacion_vendida", "Recaudación vendida ($)", 15),
    ("recaudacion_no_vendida", "Recaudación no vendida ($)", 15),
    ("origen_capacidad", "Origen de la capacidad", 34), ("fecha_scraping", "Fecha scraping", 16),
]

_PORCENTAJE = ("pct_vendido", "ocupacion_total_estimada")
_PESOS = ("recaudacion_vendida", "recaudacion_no_vendida", "recaudacion_potencial", "precio_num")


def _hoja(ws, filas, columnas):
    ws.append([c[1] for c in columnas])
    for fila in filas:
        ws.append([fila.get(c[0]) for c in columnas])
    negrita, fondo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for i, (clave, _t, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font, celda.fill = negrita, fondo
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
        formato = "0.0%" if clave in _PORCENTAJE else '"$"#,##0' if clave in _PESOS else \
            "#,##0" if clave in ("capacidad", "vendidas", "no_vendidas", "cancha_estimada", "aforo_recinto") else None
        if formato:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=i).number_format = formato
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def escribir_excel(ruta, funciones, sectores):
    wb = Workbook()
    ws = wb.active
    ws.title = "Funciones"
    _hoja(ws, funciones, COLUMNAS_FUNCIONES)
    rojo = PatternFill("solid", fgColor="F8CBAD")
    col_agotado = [c[0] for c in COLUMNAS_FUNCIONES].index("agotado") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=col_agotado).value == "SI":
            ws.cell(row=r, column=col_agotado).fill = rojo
    _hoja(wb.create_sheet("Sectores"), sectores, COLUMNAS_SECTORES)
    return _reemplazar(ruta, wb.save)


def escribir_csv(ruta, filas, columnas):
    def escribir(destino):
        with open(destino, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=[c[0] for c in columnas], extrasaction="ignore")
            w.writeheader()
            w.writerows(filas)
    return _reemplazar(ruta, escribir)


def _reemplazar(ruta, escribir):
    """Escribe en un temporal y recién después reemplaza el archivo final, para
    no dejar un reporte a medio escribir si algo falla. Si el archivo está
    abierto (Excel en Windows lo bloquea) guarda con otro nombre y avisa."""
    ruta.parent.mkdir(parents=True, exist_ok=True)
    tmp = ruta.with_name(f"~tmp_{ruta.name}")
    escribir(tmp)
    try:
        os.replace(tmp, ruta)
        return ruta
    except PermissionError:
        alternativa = ruta.with_name(f"{ruta.stem}_{datetime.now():%H%M%S}{ruta.suffix}")
        os.replace(tmp, alternativa)
        aviso(f"{ruta.name} está abierto en otro programa; guardé en {alternativa.name}")
        return alternativa
