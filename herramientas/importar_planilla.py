"""Carga al historial las lecturas de una planilla del scraper anterior.

La hoja "información general" (columnas nombre_evento, fecha_evento,
lugar_evento, sector, sector_id, asientos_totales, asientos_ocupados,
fecha_scraping) trae capacidades medidas que sirven de referencia para los
sectores que hoy están agotados.

    python herramientas/importar_planilla.py EVENTOS_AVANCE_04082026.xlsx
"""
import argparse
import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from puntoticket.capacidades import COLUMNAS_HISTORIAL  # noqa: E402
from puntoticket.config import Config  # noqa: E402
from puntoticket.fechas import parsear_fecha  # noqa: E402
from puntoticket.util import a_int  # noqa: E402


def main():
    p = argparse.ArgumentParser()
    p.add_argument("planilla")
    p.add_argument("--hoja", default="información general")
    p.add_argument("--historial", default=str(Config.historial))
    a = p.parse_args()

    ws = load_workbook(a.planilla, read_only=True, data_only=True)[a.hoja]
    filas = ws.iter_rows(values_only=True)
    enc = [str(c).strip() if c else "" for c in next(filas)]
    salida = Path(a.historial)
    nuevo = not salida.exists()
    n = 0
    with open(salida, "a", encoding="utf-8-sig" if nuevo else "utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNAS_HISTORIAL)
        if nuevo:
            w.writeheader()
        for valores in filas:
            r = dict(zip(enc, valores))
            total = a_int(r.get("asientos_totales"))
            if not total or not r.get("sector_id"):
                continue
            scrap = r.get("fecha_scraping")
            hoy = scrap.date() if isinstance(scrap, datetime) else date.today()
            fecha = r.get("fecha_evento")
            fecha = fecha.date() if isinstance(fecha, datetime) else parsear_fecha(str(fecha or ""), hoy=hoy)
            w.writerow({
                "fecha_scraping": scrap.strftime("%Y-%m-%d %H:%M") if isinstance(scrap, datetime) else str(scrap or ""),
                "recinto": r.get("lugar_evento") or "", "evento": r.get("nombre_evento") or "",
                "fecha_funcion": fecha.isoformat() if fecha else str(r.get("fecha_evento") or ""), "hora": "",
                "sector_id": r["sector_id"], "sector": r.get("sector") or "", "asientos_totales": total,
                "asientos_vendidos": a_int(r.get("asientos_ocupados")) or "",
            })
            n += 1
    print(f"{n} lecturas agregadas a {salida}")


if __name__ == "__main__":
    main()
