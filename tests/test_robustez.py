"""Escenarios de robustez: HTML cambiado, catálogo por API, Cloudflare, catálogo vacío."""
import time

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from puntoticket.config import Config
from puntoticket.scraper import correr


def _correr(servidor, tmp_path, pagina, **extra):
    cfg = Config(catalogo_url=f"{servidor}/{pagina}", salida=tmp_path / "reportes",
                 historial=tmp_path / "historial.csv", aforos=tmp_path / "no_existe.csv",
                 pausa_entre_eventos_ms=0, timeout_elementos_ms=4000, timeout_red_ms=1000,
                 timeout_asientos_ms=2500, reintentos=2, **extra)
    with sync_playwright() as p:
        nav = p.chromium.launch()
        ruta = correr(nav.new_page(), cfg)
        nav.close()
    return ruta


def _funciones(ruta):
    ws = load_workbook(ruta)["Funciones"]
    enc = [c.value for c in ws[1]]
    return [dict(zip(enc, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def test_catalogo_con_html_distinto_y_landing_sin_clases(servidor, tmp_path):
    ruta = _correr(servidor, tmp_path, "musica_v2.html")
    [f] = _funciones(ruta)  # el link del <header> no cuenta como evento
    assert f["Evento"] == "Los Bunkers" and f["Fecha"] == "2026-11-14" and f["Hora"] == "21:00"
    assert f["Lugar"] == "Teatro Caupolicán"
    assert f["Capacidad numerada (entradas)"] == 50 and f["Vendidas"] == 10


def test_catalogo_desde_json_de_la_red(servidor, tmp_path):
    [f] = _funciones(_correr(servidor, tmp_path, "musica_api.html"))
    assert f["Evento"] == "Mon Laferte" and f["Lugar"] == "Movistar Arena - Santiago Centro"
    assert f["Capacidad numerada (entradas)"] == 50


def test_espera_cloudflare_y_sigue(servidor, tmp_path):
    [f] = _funciones(_correr(servidor, tmp_path, "cloudflare.html", espera_cola_seg=20))
    assert f["Evento"] == "Los Bunkers"


def test_catalogo_vacio_no_genera_reporte(servidor, tmp_path, capsys):
    (tmp_path / "reportes").mkdir()
    previo = tmp_path / "reportes" / "puntoticket_musica_anterior.xlsx"
    previo.write_bytes(b"reporte anterior")
    assert _correr(servidor, tmp_path, "musica_vacio.html") is None
    salida = capsys.readouterr().out
    assert "ADVERTENCIA" in salida and "NO se generó reporte" in salida
    assert sorted(p.name for p in (tmp_path / "reportes").glob("*.xlsx")) == ["puntoticket_musica_anterior.xlsx"]
    assert previo.read_bytes() == b"reporte anterior"
    assert (tmp_path / "reportes" / "diagnostico" / "catalogo.html").exists()
    assert not (tmp_path / "historial.csv").exists()


def test_cola_que_no_suelta_queda_como_alerta(servidor, tmp_path):
    # La landing de Aitana manda a /queue/enqueue.html, que tarda 2,5 s: con 1 s de espera no alcanza.
    ruta = _correr(servidor, tmp_path, "musica.html", espera_cola_seg=1, filtro_titulos=["aitana"],
                   leer_asientos=False)
    log = ruta.with_name(ruta.stem + "_log.txt").read_text(encoding="utf-8")
    assert "quedó en cola" in log, log[-2000:]


def test_no_se_cuelga_con_respuesta_de_red_infinita(servidor, tmp_path):
    # Una respuesta JSON que nunca termina (analytics / long-polling) no debe trabar el scraper.
    inicio = time.time()
    ruta = _correr(servidor, tmp_path, "musica_lento.html", filtro_titulos=["anuel"])
    assert time.time() - inicio < 90
    [f] = _funciones(ruta)
    assert f["Evento"] == "Anuel AA" and f["Capacidad numerada (entradas)"] == 50


def test_no_espera_el_evento_load(servidor, tmp_path, capsys):
    # Una imagen que nunca termina de cargar retiene el evento "load" para siempre.
    inicio = time.time()
    ruta = _correr(servidor, tmp_path, "musica_sin_load.html", filtro_titulos=["anuel"])
    assert time.time() - inicio < 90
    assert [f["Evento"] for f in _funciones(ruta)] == ["Anuel AA"]
    salida = capsys.readouterr().out
    assert "Cargando catálogo" in salida and "Catálogo abierto en" in salida


def test_link_ver_mas_no_saca_del_catalogo(servidor, tmp_path, capsys):
    # Caso real: un link "Ver más" de un banner llevaba a otra página y el catálogo quedaba vacío.
    ruta = _correr(servidor, tmp_path, "musica_banner.html", filtro_titulos=["anuel"], leer_asientos=False)
    assert ruta is not None
    salida = capsys.readouterr().out
    assert "Catálogo: 2 eventos detectados" in salida


def test_reanudar_despues_de_un_corte(servidor, tmp_path, monkeypatch):
    import pytest

    from puntoticket import scraper
    original, llamados = scraper.procesar_evento, []

    def cortar_en_el_tercero(page, ev, cfg, red=None):
        if len(llamados) == 2:
            raise KeyboardInterrupt  # el usuario cierra la consola / se corta la luz
        llamados.append(ev["titulo"])
        return original(page, ev, cfg, red)

    monkeypatch.setattr(scraper, "procesar_evento", cortar_en_el_tercero)
    with pytest.raises(KeyboardInterrupt):
        _correr(servidor, tmp_path, "musica.html", leer_asientos=False)
    assert (tmp_path / "reportes" / "progreso.json").exists()

    monkeypatch.setattr(scraper, "procesar_evento", lambda page, ev, cfg, red=None:
                        llamados.append(ev["titulo"]) or original(page, ev, cfg, red))
    ruta = _correr(servidor, tmp_path, "musica.html", leer_asientos=False, reanudar=True,
                   renovar_pestana_cada=1)
    assert len(llamados) == 5 and len(set(llamados)) == 5  # ningún evento se procesó dos veces
    assert {f["Evento"] for f in _funciones(ruta)} == {"Aitana", "Karol G", "Babasónicos", "Jeff Mills", "Anuel AA"}
    assert not (tmp_path / "reportes" / "progreso.json").exists()
    assert len(list((tmp_path / "reportes").glob("*.xlsx"))) == 1  # mismo archivo, no uno nuevo
