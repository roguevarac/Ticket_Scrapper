"""Pruebas contra un sitio local que imita el HTML de PuntoTicket (tests/sitio).

Correr con:  python -m pytest -q
Usa el Chromium de Playwright (headless). No necesita internet.
"""
import csv
from datetime import date

import pytest
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from puntoticket.capacidades import COLUMNAS_HISTORIAL
from puntoticket.config import Config
from puntoticket.fechas import parsear_fecha, parsear_hora
from puntoticket.scraper import correr


@pytest.fixture(scope="module")
def reporte(servidor, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("salida")
    historial = tmp / "historial.csv"
    with open(historial, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNAS_HISTORIAL)
        w.writeheader()
        # Lectura anterior de la misma función, cuando el sector aún tenía asientos libres.
        w.writerow({"fecha_scraping": "2026-08-04 13:48", "recinto": "MOVISTAR ARENA - SANTIAGO CENTRO",
                    "evento": "Aitana", "fecha_funcion": "2026-10-19", "hora": "21:00",
                    "sector_id": "TT4S1R97", "sector": "PLATEA BAJA", "asientos_totales": 189,
                    "asientos_vendidos": 188})
        # Otro evento en el mismo recinto: solo sirve de referencia aproximada.
        w.writerow({"fecha_scraping": "2026-08-04 13:48", "recinto": "Movistar Arena - Santiago Centro",
                    "evento": "Paulo Londra", "fecha_funcion": "2026-11-06", "hora": "",
                    "sector_id": "TT4S2R92", "sector": "PLATEA ALTA GOLDEN", "asientos_totales": 474,
                    "asientos_vendidos": 100})
    aforos = tmp / "aforos.csv"
    aforos.write_text("recinto,aforo_total\nMovistar Arena - Santiago Centro,15000\n", encoding="utf-8")
    cfg = Config(catalogo_url=f"{servidor}/musica.html", salida=tmp / "reportes", historial=historial,
                 aforos=aforos, espera_cola_seg=20, pausa_entre_eventos_ms=0)
    with sync_playwright() as p:
        nav = p.chromium.launch()
        page = nav.new_page()
        ruta = correr(page, cfg)
        nav.close()
    wb = load_workbook(ruta)

    def filas(hoja):
        ws = wb[hoja]
        enc = [c.value for c in ws[1]]
        return [dict(zip(enc, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    return {"funciones": filas("Funciones"), "sectores": filas("Sectores"), "hojas": wb.sheetnames,
            "historial": historial}


def _funcion(rep, evento, fecha):
    hits = [f for f in rep["funciones"] if f["Evento"] == evento and f["Fecha"] == fecha]
    assert len(hits) == 1, (evento, fecha, [(f["Evento"], f["Fecha"], f["Hora"]) for f in rep["funciones"]])
    return hits[0]


def test_fechas_y_horas():
    assert parsear_fecha("29 de Octubre 2026") == date(2026, 10, 29)
    assert parsear_fecha("4 DE NOVIEMBRE 2026") == date(2026, 11, 4)
    assert parsear_fecha("5 de diciembre", hoy=date(2026, 10, 6)) == date(2026, 12, 5)
    assert parsear_fecha("31 de enero", hoy=date(2026, 10, 6)) == date(2027, 1, 31)
    assert parsear_fecha("05-12-2026") == date(2026, 12, 5)
    assert parsear_hora("Lunes 19 de octubre 2026 - 21:00 hrs") == "21:00"
    assert parsear_hora("19 de octubre 2026") == ""
    assert parsear_hora("Entradas desde $19.000") == ""
    assert parsear_hora("Show 9 pm") == "21:00"


def test_cada_fecha_es_una_funcion(reporte):
    claves = {(f["Evento"], f["Fecha"]) for f in reporte["funciones"]}
    assert claves == {
        ("Aitana", "2026-10-19"), ("Aitana", "2026-10-20"),
        ("Karol G", "2026-12-10"), ("Karol G", "2026-12-11"),
        ("Babasónicos", "2026-11-08"), ("Jeff Mills", "2026-12-05"), ("Anuel AA", "2026-11-20"),
    }


def test_aitana_ocupacion_y_capacidad_de_agotados(reporte):
    f = _funcion(reporte, "Aitana", "2026-10-19")
    assert (f["Día"], f["Hora"], f["Estado"], f["Agotado"]) == ("lunes", "21:00", "DISPONIBLE", "NO")
    assert f["Lugar"] == "Movistar Arena - Santiago Centro"
    # Medidos: 90/30, 288/288, 100/95. Agotados: TT4S1R97=189 (misma función), TT4S2R92=474 (recinto).
    assert f["Capacidad numerada (entradas)"] == 90 + 288 + 100 + 189 + 474
    assert f["Vendidas"] == 30 + 288 + 95 + 189 + 474
    assert f["Sectores numerados"] == 5 and f["Sectores agotados"] == 3
    assert f["Sectores sin asiento (cancha/general)"] == "CANCHA GENERAL (DISPONIBLE)"
    assert f["Cancha estimada"] == 15000 - 1141
    assert f["Calidad del dato"].startswith("Parcial")  # usó una referencia de otro evento
    sec = {s["Sector ID"]: s for s in reporte["sectores"] if s["Evento"] == "Aitana" and s["Fecha"] == "2026-10-19"}
    assert sec["TT2S1"]["Sector"] == "DIAMANTE" and sec["TT2S1"]["Vendidas"] == 30
    assert sec["TT4S1R97"]["Sector"] == "PLATEA BAJA" and sec["TT4S1R97"]["Origen de la capacidad"] == "Lectura anterior (misma función)"
    assert sec["TT4S2R92"]["Origen de la capacidad"].startswith("Estimado")
    assert sec["TT6S1R-40"]["Sector"] == "TRIBUNA" and sec["TT6S1R-40"]["Capacidad (entradas)"] == 100
    assert sec["TT5"]["Tipo"] == "GENERAL"
    # Recaudación a precio de lista: DIAMANTE $120.000, PLATEA BAJA $86.250, TRIBUNA $11.500.
    assert sec["TT2S1"]["Precio ($)"] == 120000 and sec["TT2S1"]["No vendidas"] == 60
    assert f["Recaudación vendida ($, precio lista)"] == (30 + 288) * 120000 + (189 + 474) * 86250 + 95 * 11500
    assert f["Recaudación no vendida ($, precio lista)"] == 60 * 120000 + 5 * 11500
    assert f["No vendidas"] == 60 + 5


def test_porcentaje_pondera_por_entradas(reporte):
    # 5 sectores: el % no es "sectores vendidos" sino entradas vendidas / entradas totales.
    f = _funcion(reporte, "Aitana", "2026-10-19")
    assert abs(f["% vendido (por entradas)"] - 1076 / 1141) < 1e-9


def test_reporte_sin_columnas_ni_hojas_descartadas(reporte):
    f = reporte["funciones"][0]
    for col in ("Género", "Fecha (texto sitio)", "Alertas", "Confianza"):
        assert col not in f
    s = reporte["sectores"][0]
    assert "Estado lista precios" not in s and "URL compra" not in s
    assert reporte["hojas"] == ["Funciones", "Sectores"]


def test_funcion_agotada_sin_link(reporte):
    f = _funcion(reporte, "Aitana", "2026-10-20")
    assert f["Agotado"] == "SI" and f["% vendido (por entradas)"] == 1.0 and f["Día"] == "martes"


def test_karolg_fechas_desde_imagen_y_dos_botones(reporte):
    f = _funcion(reporte, "Karol G", "2026-12-10")
    assert f["Hora"] == "20:30" and f["Agotado"] == "NO"
    # Los dos botones de la misma fecha se juntan en una función: cancha, palco y pacífico.
    assert f["Sectores numerados"] == 2 and f["Capacidad numerada (entradas)"] == 200
    assert "CANCHA (AGOTADO)" in f["Sectores sin asiento (cancha/general)"]
    assert _funcion(reporte, "Karol G", "2026-12-11")["Estado"] == "SIN INFORMACIÓN"


def test_landing_agotada_pero_compra_disponible(reporte):
    f = _funcion(reporte, "Anuel AA", "2026-11-20")
    assert f["Estado en landing"] == "AGOTADO"
    assert f["Agotado"] == "NO" and f["Vendidas"] == 10 and f["Capacidad numerada (entradas)"] == 50
    assert f["Calidad del dato"] == "Todo medido hoy"


def test_agotado_en_catalogo_y_template_address(reporte):
    b = _funcion(reporte, "Babasónicos", "2026-11-08")
    assert b["Agotado"] == "SI" and b["% ocupación total (estim.)"] == 1.0
    j = _funcion(reporte, "Jeff Mills", "2026-12-05")
    assert j["Hora"] == "23:30" and j["Lugar"] == "Club Chocolate"


def test_historial_acumula_lecturas_medidas(reporte):
    with open(reporte["historial"], encoding="utf-8-sig") as f:
        filas = list(csv.DictReader(f))
    ids = {(r["evento"], r["sector_id"]) for r in filas}
    assert ("Aitana", "TT2S1") in ids and ("Anuel AA", "TT2S1R97") in ids
